import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding="utf-8")
path = r"c:\Users\safee\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm\LocalState\sessions\885E58C0E755C4842673FC1E48B7F6610BCD37C7\transfers\2026-38\Mini project Review 1.pptx"
prs = Presentation(path)
print("slides", len(prs.slides))
for i, slide in enumerate(prs.slides, 1):
    print("\n" + "=" * 70)
    print(f"SLIDE {i}")
    for shape in slide.shapes:
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                print(t)
                print("---")
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            for row in shape.table.rows:
                print(" | ".join(c.text.strip().replace("\n", " ") for c in row.cells))
            print("---TABLE---")
