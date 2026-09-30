"""Patch Review-1 PPTX: content + UML images, then copy to Desktop."""
from copy import deepcopy
from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.oxml.ns import qn

ROOT = Path(r"c:\Users\safee\OneDrive\Desktop\Mini Project")
SRC = ROOT / "MediMesh_Mini_Project_Review1.pptx"
IMG = ROOT / "scripts" / "pptx_extract"
OUT_PROJECT = ROOT / "MediMesh_Mini_Project_Review1.pptx"
OUT_DESKTOP = Path(r"c:\Users\safee\OneDrive\Desktop\MediMesh_Mini_Project_Review1.pptx")
WHATSAPP = Path(
    r"c:\Users\safee\AppData\Local\Packages\5319275A.WhatsAppDesktop_cv1g1gvanyjgm"
    r"\LocalState\sessions\885E58C0E755C4842673FC1E48B7F6610BCD37C7"
    r"\transfers\2026-37\MediMesh_Mini_Project_Review1.pptx"
)


def set_runs(paragraph, text):
    if paragraph.runs:
        paragraph.runs[0].text = text
        for run in paragraph.runs[1:]:
            run.text = ""
    else:
        paragraph.text = text


def write_lines(shape, lines):
    tf = shape.text_frame
    while len(tf.paragraphs) < len(lines):
        last = tf.paragraphs[-1]._p
        last.addnext(deepcopy(last))
    for i, item in enumerate(lines):
        level, text = item if isinstance(item, tuple) else (0, item)
        p = tf.paragraphs[i]
        p.level = level
        set_runs(p, text)
    for i in range(len(lines), len(tf.paragraphs)):
        set_runs(tf.paragraphs[i], "")


def write_table(shape, rows):
    table = shape.table
    for r, row in enumerate(rows):
        for c, val in enumerate(row):
            if r >= len(table.rows) or c >= len(table.columns):
                continue
            set_runs(table.cell(r, c).text_frame.paragraphs[0], val)


def replace_picture(shape, path):
    blip = shape._element.find(".//" + qn("a:blip"))
    if blip is None:
        raise RuntimeError("no blip")
    r_id = blip.get(qn("r:embed"))
    part = shape.part.related_part(r_id)
    part.blob = Path(path).read_bytes()


def largest_picture(slide):
    pics = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE]
    if not pics:
        return None
    return max(pics, key=lambda s: s.width * s.height)


def content_shape(slide):
    for s in slide.shapes:
        if s.has_text_frame and "Content Placeholder" in s.name:
            return s
    return None


def title_named(slide, name_part="Title"):
    for s in slide.shapes:
        if s.has_text_frame and name_part in s.name and s.text_frame.text.strip():
            return s
    return None


prs = Presentation(str(SRC))

# Slide 1 title
for s in prs.slides[0].shapes:
    if s.has_text_frame and "Title of the Project" in s.text_frame.text:
        set_runs(s.text_frame.paragraphs[0], "Title: MediMesh AI — Explainable Multi-Hospital Care")

# Slide 2 outlines
write_lines(content_shape(prs.slides[1]), [
    "Main idea of the Project",
    "Existing System Vs Proposed",
    "Proposed Project Objectives",
    "Proposed Project Outcomes",
    "Software and Hardware Requirements",
    "Process Model",
    "Architectural Diagram [Modified]",
    "UML Diagrams — Use Case, Class, Sequence, Activity, State Chart",
    "Proposed Methods and Algorithms [with working demo example]",
])

# Slide 3 main idea
write_lines(content_shape(prs.slides[2]), [
    "The problem",
    (1, "Smaller hospitals still run booking, waiting, prescriptions and stock on disconnected registers."),
    (1, "Patients cannot see the wait maths; reports are jargon; night staff is not who the app suggests."),
    "The solution: MediMesh AI",
    (1, "One affordable app for patient, doctor, nurse, pharmacist, receptionist and administrator."),
    (1, "Three isolated hospitals (Sunrise, Lotus, Arogya). Stock never leaks. Only anonymised demand is pooled."),
    "What is unique (not a generic HMS clone)",
    (1, "Live OPD token on Confirm Appointment + dual wait: your turn vs hospital-wide HospitalFlow formula."),
    (1, "Care Copilot routes by who is on this shift (day OPD / night nurse). It is not a diagnosis."),
    (1, "Care Explainer in English / Hindi / Telugu + read-aloud, only after the doctor approves."),
])

# Slide 4 comparison table
for s in prs.slides[3].shapes:
    if s.shape_type == MSO_SHAPE_TYPE.TABLE:
        write_table(s, [
            ["Aspect", "Existing system", "MediMesh (proposed)"],
            ["Booking", "Phone / walk-in registers, no slot visibility", "Calendar booking; Confirm mints a live OPD token"],
            ["Waiting", "Unknown, blind wait", "Dual wait: your turn maths vs hospital-wide HospitalFlow"],
            ["Routing", "Patient guesses the clinic", "Care Copilot: right queue for this shift; emergency is not booked"],
            ["Medical info", "Jargon-heavy reports", "Doctor-approved explainer EN / HI / TE + read-aloud"],
            ["Handover", "Verbal / paper, lossy", "Signed day / night handover on the nurse dashboard"],
            ["Pharmacy", "Manual re-entry; stock guessed", "Digital Rx; in / low / out; mesh demand without sharing stock"],
            ["Access", "Shared logins, no hospital wall", "Role + hospital_id isolation (staff and stock stay inside one hospital)"],
            ["Cost / scale", "Built for large hospital chains", "Affordable 3-hospital mesh for smaller facilities"],
        ])
    if s.has_text_frame and "DEPARTMENT OF CSE" in s.text_frame.text:
        set_runs(s.text_frame.paragraphs[0], "DEPARTMENT OF CSD MINI PROJECT REVIEW-1")

# Slide 5 objectives
write_lines(content_shape(prs.slides[4]), [
    "Unify booking and live doctor availability across three hospitals, with hospital-level isolation.",
    "Mint a live OPD token on confirm and show dual wait: people ahead × consult ÷ doctors vs HospitalFlow total.",
    "Explain waits with a visible formula (queue, pharmacy, stock-outs, staffing, handover, crowd) — not a black box.",
    "Route symptoms with Care Copilot to the staff on this shift; refuse diagnosis; divert chest tightness to casualty.",
    "Simplify treatment in EN / HI / TE with read-aloud after the doctor approves the explainer.",
    "Digitize prescription-to-pharmacy (in / low / out) and pool anonymised SKU demand for cheaper mesh buying.",
    "Support day / night nurse handover that is written, signed and visible to the next shift.",
    "Give administrators staffing sliders that recalculate HospitalFlow live, on a low-cost stack.",
])

# Slide 6 outcomes
write_lines(content_shape(prs.slides[5]), [
    "Patient sees a token and a wait they can argue with (your turn vs hospital-wide), like a Rapido ETA.",
    "Same symptoms give different next steps by shift — night nurse now, GP when clinic opens.",
    "Patients understand Metformin / recovery in their language; doctors stay accountable via approval.",
    "Pharmacy knows in / low / out; mesh demand flags what three hospitals should buy together.",
    "Handover is not lost between Kavitha (day) and Imran (night).",
    "Queues, stock and files do not leak across Sunrise, Lotus and Arogya.",
    "Admin can change doctors-on-duty and watch the formula move in the viva demo.",
    "A reusable, affordable coordination template — not another Practo / HMS clone.",
])

# Slide 7 software table
for s in prs.slides[6].shapes:
    if s.shape_type == MSO_SHAPE_TYPE.TABLE:
        write_table(s, [
            ["Category", "Component", "Choice"],
            ["Software", "Frontend", "React + TypeScript + Vite + Tailwind (phone-shell UI)"],
            ["Software", "State / roles", "MeshContext reducer; session picks hospital then role (demo login)"],
            ["Software", "Analytics AI", "Python FastAPI HospitalFlow; same formula also runs on-device in TS"],
            ["Software", "Database", "In-memory seed for demo; supabase/schema.sql with RLS by hospital_id"],
            ["Software", "Care Copilot", "Rule-based, shift-aware routing (not an LLM diagnosis)"],
            ["Software", "Explainer", "Doctor-approved cards + SpeechSynthesis (EN / HI / TE)"],
            ["Hardware", "Dev machine", "i5 / Ryzen 5, 8GB+ RAM, Node 20, Python 3.11"],
            ["Hardware", "Server (pilot)", "Local Vite :5173 + uvicorn :8000; cloud VM later"],
            ["Hardware", "Client devices", "Smartphone or desktop browser"],
            ["Hardware", "Pharmacy", "Optional barcode scanner later; demo uses tap to dispense"],
        ])

# Slide 8 process model
write_lines(content_shape(prs.slides[7]), [
    "Recommended model: Incremental / Agile",
    (1, "Each review shows a working demo. Requirements stay allowed to change. Fits a small team."),
    "Increment 1 — Multi-hospital login, role dashboards, calendar booking",
    "Increment 2 — Care Explainer (EN/HI/TE + voice) and Care Copilot (shift-aware routing)",
    "Increment 3 — Live token on confirm + dual wait + HospitalFlow FastAPI",
    "Increment 4 — Pharmacy in/low/out, mesh buying power, admin staffing sliders",
])

# Diagrams
mapping = {
    9: IMG / "uml_architecture.png",
    10: IMG / "uml_usecase.png",
    11: IMG / "uml_class.png",
    12: IMG / "uml_sequence.png",
    13: IMG / "uml_activity.png",
    14: IMG / "uml_state.png",
}
for num, path in mapping.items():
    pic = largest_picture(prs.slides[num - 1])
    if pic is None:
        raise RuntimeError(f"no picture on slide {num}")
    replace_picture(pic, path)
    print("replaced slide", num, path.name)

# Slide 15 algorithms
algo = content_shape(prs.slides[14]) or prs.slides[14].shapes[3]
write_lines(algo, [
    "HospitalFlow (visible wait formula)",
    (1, "consult = waiting × avgConsult ÷ max(doctors,1); then + pharmacy + stock-outs×6; × staff × handover × crowd."),
    (1, "Demo: 3 people × 12 min ÷ 4 doctors = 9 min YOUR TURN. Hospital-wide ~30 min includes pharmacy and crowd."),
    "Care Copilot (shift-aware routing, not diagnosis)",
    (1, "Match words to a specialty, then to who is on duty. Night → night nurse + hold morning slot. Chest tightness → casualty."),
    (1, "Example: “thirsty and tired” + Day → Dr. Meera Iyer (GP) + Metformin. Same words + Night → Imran Shaik, clinic 08:00."),
    "Mesh pharmacy intelligence (no stock sharing)",
    (1, "Each hospital keeps its own quantity. Mesh only sums soldWeek by SKU and lists which hospitals are OUT."),
    (1, "Example: Arogya Amlodipine quantity=0; Sunrise still has stock — Arogya cannot see Sunrise shelves."),
])

prs.save(str(OUT_PROJECT))
prs.save(str(OUT_DESKTOP))
if WHATSAPP.parent.exists():
    try:
        prs.save(str(WHATSAPP))
        print("also saved WhatsApp copy")
    except OSError as e:
        print("WhatsApp copy skipped:", e)
print("saved", OUT_PROJECT)
print("saved", OUT_DESKTOP)
print("bytes", OUT_DESKTOP.stat().st_size)
