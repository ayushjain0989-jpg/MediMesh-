import json
import sys
from pathlib import Path
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding="utf-8")
root = Path(r"c:\Users\safee\OneDrive\Desktop\Mini Project")
prs = Presentation(str(root / "MediMesh_Mini_Project_Review1.pptx"))
out = root / "scripts" / "pptx_extract"
out.mkdir(parents=True, exist_ok=True)

def dump_table(shape):
    rows = []
    for row in shape.table.rows:
        rows.append([cell.text.strip().replace("\n", " ") for cell in row.cells])
    return rows

for i, slide in enumerate(prs.slides, 1):
    print("\n" + "=" * 60)
    print(f"SLIDE {i}")
    img_n = 0
    for j, shape in enumerate(slide.shapes):
        if shape.has_text_frame:
            t = shape.text_frame.text.strip()
            if t:
                print(f"  text[{j}] {t[:1000]}")
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            print("  TABLE:")
            for row in dump_table(shape):
                print("   | " + " | ".join(row))
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
            img_n += 1
            blob = shape.image.blob
            ext = shape.image.ext
            dest = out / f"slide{i}_pic{img_n}.{ext}"
            dest.write_bytes(blob)
            print(f"  IMAGE[{img_n}] {dest.name} {len(blob)} bytes {shape.width}x{shape.height}")
