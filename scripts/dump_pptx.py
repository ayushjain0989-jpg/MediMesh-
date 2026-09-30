from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pathlib import Path

p = Path(r"c:\Users\safee\OneDrive\Desktop\Mini Project\MediMesh_Mini_Project_Review1.pptx")
prs = Presentation(str(p))
print("slides", len(prs.slides))
print("w", prs.slide_width, "h", prs.slide_height)
for i, slide in enumerate(prs.slides, 1):
    print("\n" + "=" * 60)
    print(f"SLIDE {i}")
    if slide.has_notes_slide and slide.notes_slide.notes_text_frame:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if notes:
            print("NOTES:", notes[:400])
    for j, shape in enumerate(slide.shapes):
        st = shape.shape_type
        extra = ""
        if st == MSO_SHAPE_TYPE.PICTURE:
            extra = f" PICTURE {shape.image.content_type}"
        elif st == MSO_SHAPE_TYPE.GROUP:
            extra = " GROUP"
        elif st == MSO_SHAPE_TYPE.TABLE:
            extra = " TABLE"
        text = ""
        if shape.has_text_frame:
            text = shape.text_frame.text.strip().replace("\n", " | ")
        print(f"  [{j}] {st} {shape.name}{extra}")
        if text:
            print(f"      TEXT: {text[:800]}")
        if st == MSO_SHAPE_TYPE.GROUP:
            for k, child in enumerate(shape.shapes):
                ct = child.text_frame.text.strip().replace("\n", " | ") if child.has_text_frame else ""
                print(f"      child[{k}] {child.shape_type} {child.name} {ct[:200]}")
