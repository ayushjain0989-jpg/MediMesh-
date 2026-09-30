import sys
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE

sys.stdout.reconfigure(encoding="utf-8")
prs = Presentation(r"c:\Users\safee\OneDrive\Desktop\Mini Project\MediMesh_Mini_Project_Review1.pptx")

for i in [1, 3, 4, 5, 6, 7, 8, 15]:
    slide = prs.slides[i - 1]
    print(f"\n===== SLIDE {i} =====")
    for j, shape in enumerate(slide.shapes):
        if not shape.has_text_frame:
            continue
        print(f" shape[{j}] {shape.name}")
        for pi, p in enumerate(shape.text_frame.paragraphs):
            font = p.font
            size = font.size.pt if font.size else None
            bold = font.bold
            name = font.name
            print(f"  p{pi} lvl={p.level} bold={bold} size={size} name={name} | {p.text[:180]!r}")
        if shape.shape_type == MSO_SHAPE_TYPE.TABLE:
            for r, row in enumerate(shape.table.rows):
                for c, cell in enumerate(row.cells):
                    t = cell.text_frame.paragraphs[0]
                    size = t.font.size.pt if t.font.size else None
                    print(f"  cell[{r},{c}] size={size} bold={t.font.bold} | {cell.text[:80]!r}")
